"""Phase 4 Step 4: LLM-as-judge, cross-judged and validated by hand (docs/eval_plan.md, section 4
and post-hoc change 2).

Commands:
  run --judge groq | groq-qwen   the judge grades the OTHER LLM's fixed Step 3 answers (in-scope,
                                 answered). Input: question, gold chunk text (with status), reference
                                 answer, the sources the answer cited (with status), the answer. Never
                                 the full context. Cached per (answer key, answer text, judge model,
                                 rubric version); stops cleanly on quota / spend-limit / API errors.
  export-manual                  20 stratified answers (10 per answering LLM, slices in proportion,
                                 seed 42), blind (no model name, no judge labels) -> manual_labels.xlsx
  report                         judge label distributions per answering LLM; once the workbook is
                                 filled in, % agreement and Cohen's kappa per criterion and the plan's
                                 kappa rule -> reports/eval_judge.json and .md
"""

from __future__ import annotations

import argparse
import hashlib
import json
import logging
import random
import re
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from maritime_rag.answer.providers import Generation, Message, ProviderError, make_provider
from maritime_rag.config import resolve
from maritime_rag.eval.answers import QUOTA_HINTS, Setup, cache_path, current_records, load_contexts, read_cache, setup
from maritime_rag.eval.questions import OUT_OF_SCOPE, GoldText

CRITERIA = ("correctness", "faithfulness", "superseded_mentioned")
_NA = {"n.a.", "n.a", "na", "n/a", "not applicable", "none"}


def rubric_version(jcfg: dict[str, Any]) -> str:
    blob = json.dumps({"prompt": jcfg["prompt"], "labels": jcfg["labels"]}, sort_keys=True, ensure_ascii=False)
    return hashlib.sha256(blob.encode()).hexdigest()[:10]


def text_hash(text: str) -> str:
    return hashlib.sha256((text or "").encode()).hexdigest()[:10]


def judged_answers(s: Setup, answerer: str, data: dict[str, Any]) -> list[dict[str, Any]]:
    """The fixed Step 3 answers that are judged: in-scope and answered."""
    records, _ = current_records(s, answerer, data)
    return [r for r in records if r["slice"] != OUT_OF_SCOPE and not r["refused"]]


def render_sources(gold: GoldText, items: list[dict[str, Any]] | list[str]) -> str:
    """[id | location | pages | status] + text for each chunk id (status from the version index)."""
    ids = [i["chunk_id"] if isinstance(i, dict) else i for i in items]
    return gold.render(ids) if ids else "(none)"


def judge_messages(s: Setup, gold: GoldText, answer: dict[str, Any]) -> list[Message]:
    q = next(q for q in s.questions if q["id"] == answer["id"])
    user = "\n\n".join([
        f"QUESTION:\n{q['question']}",
        f"REFERENCE ANSWER:\n{q['reference_answer']}",
        f"GOLD SOURCES:\n{render_sources(gold, q['gold_chunk_ids'])}",
        f"CITED SOURCES (only these were cited by the answer):\n{render_sources(gold, answer['cited'])}",
        f"ANSWER:\n{answer['text']}",
    ])
    return [{"role": "system", "content": s.ecfg["judge"]["prompt"]["system"]}, {"role": "user", "content": user}]


def parse_judgement(text: str, labels: dict[str, list[str]]) -> tuple[dict[str, str] | None, str | None]:
    """The judge's JSON object -> labels (or None, error). Tolerates code fences and extra text."""
    m = re.search(r"\{.*\}", text or "", re.DOTALL)
    if not m:
        return None, "no JSON object in the judge output"
    try:
        raw = json.loads(m[0])
    except json.JSONDecodeError as e:
        return None, f"invalid JSON: {e}"
    if not isinstance(raw, dict):
        return None, "JSON is not an object"
    got = {str(k).strip().lower(): v for k, v in raw.items()}
    out: dict[str, str] = {}
    for crit in CRITERIA:
        v = str(got.get(crit, "")).strip().lower()
        if crit == "superseded_mentioned" and v in _NA:
            v = "n.a."
        if v not in labels[crit]:
            return None, f"{crit}: {v!r} is not one of {labels[crit]}"
        out[crit] = v
    out["rationale"] = str(got.get("rationale", "")).strip()
    return out, None


def judge_cache_path(s: Setup, judge: str, answerer: str) -> Path:
    return cache_path(s, judge).with_name(f"judge_{judge}_on_{answerer}.jsonl")


def judge_key(answer: dict[str, Any], judge_model: str, rv: str) -> str:
    return f"{answer['key']}|{text_hash(answer['text'])}|{judge_model}|{rv}"


def _gen_dict(g: Generation | None) -> dict[str, Any] | None:
    if g is None:
        return None
    return {"model": g.model, "prompt_tokens": g.prompt_tokens, "output_tokens": g.output_tokens,
            "reasoning_tokens": g.reasoning_tokens, "latency_s": g.latency_s, "wait_s": g.wait_s,
            "retries": g.retries, "finish_reason": g.finish_reason, "cost_usd": g.cost_usd}


def run_judge(s: Setup, judge: str, provider: Any = None, out: Any = sys.stdout) -> int:
    """Grade every uncached answer of the judged LLM. Returns the number of new judgements."""
    jcfg = s.ecfg["judge"]
    answerer = jcfg["pairs"][judge]
    data = load_contexts(s)
    answers = judged_answers(s, answerer, data)
    model = s.acfg["llm"]["profiles"][judge]["model"]
    rv = rubric_version(jcfg)
    path = judge_cache_path(s, judge, answerer)
    done = {r["key"] for r in read_cache(path)}
    todo = [a for a in answers if judge_key(a, model, rv) not in done]
    print(f"{judge} ({model}) judges {answerer}: {len(answers) - len(todo)} cached, {len(todo)} to judge "
          f"(rubric version {rv})", file=out)
    if not todo:
        return 0
    if provider is None:
        provider = make_provider(s.acfg["llm"], judge, calls_file=resolve(s.acfg["logging"]["calls_file"]))
        provider.check()
    gold = GoldText(s.corpus)
    path.parent.mkdir(parents=True, exist_ok=True)
    new = 0
    try:
        for a in todo:
            msgs = judge_messages(s, gold, a)
            gen = provider.generate(msgs)
            labels, err = parse_judgement(gen.text, jcfg["labels"])
            gens = [_gen_dict(gen)]
            if labels is None:  # one retry with a format reminder
                retry = msgs + [{"role": "assistant", "content": gen.text},
                                {"role": "user", "content": "Reply with only the JSON object described above."}]
                gen = provider.generate(retry)
                labels, err = parse_judgement(gen.text, jcfg["labels"])
                gens.append(_gen_dict(gen))
            row = {"key": judge_key(a, model, rv), "answer_key": a["key"], "answer_text_hash": text_hash(a["text"]),
                   "id": a["id"], "slice": a["slice"], "answerer": answerer, "judge": judge, "judge_model": model,
                   "rubric_version": rv, "labels": labels, "error": err, "raw": gen.text,
                   "generations": gens, "time": datetime.now().isoformat(timespec="seconds")}
            with path.open("a", encoding="utf-8") as f:
                f.write(json.dumps(row, ensure_ascii=False) + "\n")
            new += 1
            shown = " / ".join(labels[c] for c in CRITERIA) if labels else f"INVALID ({err})"
            print(f"  {a['id']}: {shown}", file=out)
    except ProviderError as e:
        kind = "quota / spend limit" if any(h in str(e).lower() for h in QUOTA_HINTS) else "provider error"
        print(f"stopped ({kind}) after {new} new judgements: {e}\nre-run the same command to resume", file=out)
    except KeyboardInterrupt:
        print(f"interrupted after {new} new judgements; re-run to resume", file=out)
    finally:
        total = sum(g.get("cost_usd") or 0.0 for r in read_cache(path) for g in r["generations"] if g)
        print(f"this run: {getattr(provider, 'calls', 0)} calls (incl. probe), {getattr(provider, 'prompt_tokens', 0)} "
              f"prompt + {getattr(provider, 'output_tokens', 0)} output tokens, est. ${getattr(provider, 'cost_usd', 0.0):.4f}; "
              f"all cached judgements of {answerer} by {judge}: est. ${total:.4f}", file=out)
    return new


def current_judgements(s: Setup, judge: str, data: dict[str, Any]) -> dict[str, dict[str, Any]]:
    """answer key -> latest judgement for the current answer text, judge model and rubric version."""
    jcfg = s.ecfg["judge"]
    answerer = jcfg["pairs"][judge]
    model, rv = s.acfg["llm"]["profiles"][judge]["model"], rubric_version(jcfg)
    want = {judge_key(a, model, rv): a["key"] for a in judged_answers(s, answerer, data)}
    out = {}
    for r in read_cache(judge_cache_path(s, judge, answerer)):
        if r["key"] in want:
            out[want[r["key"]]] = r
    return out


# --- manual sample -------------------------------------------------------------------------------

def largest_remainder(counts: dict[str, int], total: int) -> dict[str, int]:
    """Allocate `total` across slices in proportion to `counts` (largest remainder, ties by name)."""
    n = sum(counts.values())
    if n == 0:
        return {k: 0 for k in counts}
    quotas = {k: total * v / n for k, v in counts.items()}
    alloc = {k: int(q) for k, q in quotas.items()}
    rest = sorted(counts, key=lambda k: (-(quotas[k] - alloc[k]), k))
    for k in rest[: total - sum(alloc.values())]:
        alloc[k] += 1
    return {k: min(alloc[k], counts[k]) for k in counts}


def stratified_sample(s: Setup, data: dict[str, Any]) -> list[dict[str, Any]]:
    """per_llm answers per answering LLM, slices in proportion; blind ids M01.. in shuffled order."""
    m = s.ecfg["judge"]["manual"]
    seed = s.ecfg["seed"]
    picked = []
    for llm in s.ecfg["answers"]["llms"]:
        pool = judged_answers(s, llm, data)
        by_slice: dict[str, list[dict[str, Any]]] = {}
        for a in pool:
            by_slice.setdefault(a["slice"], []).append(a)
        alloc = largest_remainder({k: len(v) for k, v in by_slice.items()}, m["per_llm"])
        for sl in sorted(by_slice):
            rng = random.Random(f"{seed}-{llm}-{sl}")
            for a in rng.sample(sorted(by_slice[sl], key=lambda r: r["id"]), alloc[sl]):
                picked.append({"llm": llm, "answer": a})
    random.Random(seed).shuffle(picked)
    return [{"sample_id": f"M{i:02d}", "llm": p["llm"], "id": p["answer"]["id"], "slice": p["answer"]["slice"],
             "answer_key": p["answer"]["key"], "answer_text_hash": text_hash(p["answer"]["text"]),
             "answer": p["answer"]} for i, p in enumerate(picked, start=1)]


MANUAL_COLUMNS = ["sample_id", "slice", "question", "reference_answer", "gold sources", "cited sources", "answer",
                  "correctness", "faithfulness", "superseded_mentioned", "notes"]


def export_manual(s: Setup, data: dict[str, Any], xlsx: Path, mapping: Path) -> list[dict[str, Any]]:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    sample = stratified_sample(s, data)
    gold = GoldText(s.corpus)
    labels = s.ecfg["judge"]["labels"]
    qs = {q["id"]: q for q in s.questions}
    wb = Workbook()
    ws = wb.active
    ws.title = "labels"
    ws.append(MANUAL_COLUMNS)
    for x in sample:
        q, a = qs[x["id"]], x["answer"]
        ws.append([x["sample_id"], x["slice"], q["question"], q["reference_answer"],
                   render_sources(gold, q["gold_chunk_ids"]), render_sources(gold, a["cited"]), a["text"], "", "", "", ""])
    widths = {"sample_id": 9, "slice": 12, "question": 40, "reference_answer": 40, "gold sources": 60,
              "cited sources": 60, "answer": 70, "correctness": 14, "faithfulness": 14, "superseded_mentioned": 14,
              "notes": 30}
    for i, name in enumerate(MANUAL_COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        ws.column_dimensions[cell.column_letter].width = widths[name]
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    last = len(sample) + 1
    for crit in ("correctness", "faithfulness", "superseded_mentioned"):
        col = ws.cell(row=1, column=MANUAL_COLUMNS.index(crit) + 1).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(labels[crit]) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")
        for r in range(2, last + 1):
            ws[f"{col}{r}"].fill = PatternFill("solid", fgColor="FFF2CC")
    ws.freeze_panes = "C2"
    readme = wb.create_sheet("README")
    for line in [
        "Manual labels for the judge validation (docs/eval_plan.md section 4, post-hoc change 2).",
        "The answering model and the judge's labels are hidden. Fill all three dropdowns on every row.",
        "",
        "correctness, against the gold sources and the reference answer:",
        "  correct = all key points of the reference, nothing contradicting the gold text",
        "  partial = some key points missing, or a minor error that does not change the main answer",
        "  wrong = contradicts the gold text, answers a different question, or the main point is missing",
        "  (an answer from a different provision that states the same rule can be correct)",
        "faithfulness, against the CITED sources only (not the gold):",
        "  supported = every statement follows from the sources the answer cites",
        "  partly = at least one statement is not in the cited sources (added qualifier, wrong party, inference)",
        "  unsupported = the main claims are not in the cited sources",
        "superseded_mentioned: n.a. if no gold or cited source relevant to the question is SUPERSEDED/REPEALED;",
        "  otherwise yes if the answer says the old version was amended/repealed (and by which document), else no",
        "Ignore the fixed 'not legal advice' line. When done, save and run: uv run maritime-eval-judge report",
    ]:
        readme.append([line])
    readme.column_dimensions["A"].width = 110
    xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx)
    mapping.write_text(json.dumps([{k: v for k, v in x.items() if k != "answer"} for x in sample], indent=1),
                       encoding="utf-8")
    return sample


def read_manual(xlsx: Path, labels: dict[str, list[str]]) -> tuple[dict[str, dict[str, str]], list[str]]:
    """sample id -> {criterion: label}; errors for empty or invalid cells."""
    from openpyxl import load_workbook

    ws = load_workbook(xlsx, read_only=True, data_only=True)["labels"]
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h) for h in rows[0]]
    out, errors = {}, []
    for r in rows[1:]:
        if not r or not r[head.index("sample_id")]:
            continue
        sid = str(r[head.index("sample_id")])
        got = {}
        for crit in CRITERIA:
            v = str(r[head.index(crit)] or "").strip().lower()
            if crit == "superseded_mentioned" and v in _NA:
                v = "n.a."
            if v not in labels[crit]:
                errors.append(f"{sid}: {crit} is {v or 'empty'!r}")
            got[crit] = v
        out[sid] = got
    return out, errors


# --- agreement and report ------------------------------------------------------------------------

def cohen_kappa(a: list[str], b: list[str]) -> float | None:
    """Unweighted Cohen's kappa of two label lists (None if empty)."""
    n = len(a)
    if n == 0:
        return None
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[k] * cb[k] for k in set(a) | set(b)) / (n * n)
    if pe == 1.0:
        return 1.0 if po == 1.0 else 0.0
    return (po - pe) / (1 - pe)


def agreement(s: Setup, data: dict[str, Any], manual: dict[str, dict[str, str]],
              mapping: list[dict[str, Any]]) -> dict[str, Any]:
    """Manual vs judge labels per criterion. The judge of an answer is the other LLM."""
    jcfg = s.ecfg["judge"]
    judge_of = {answerer: judge for judge, answerer in jcfg["pairs"].items()}
    judgements = {judge: current_judgements(s, judge, data) for judge in jcfg["pairs"]}
    pairs: dict[str, tuple[list[str], list[str]]] = {c: ([], []) for c in CRITERIA}
    missing = []
    for x in mapping:
        j = judgements[judge_of[x["llm"]]].get(x["answer_key"])
        if j is None or j["labels"] is None or j["answer_text_hash"] != x["answer_text_hash"]:
            missing.append(x["sample_id"])
            continue
        for c in CRITERIA:
            pairs[c][0].append(manual[x["sample_id"]][c])
            pairs[c][1].append(j["labels"][c])
    out: dict[str, Any] = {"n": len(mapping) - len(missing), "missing_judgements": missing, "criteria": {}}
    for c, (m, j) in pairs.items():
        kappa = cohen_kappa(m, j)
        out["criteria"][c] = {
            "agreement": sum(x == y for x, y in zip(m, j)) / len(m) if m else None, "kappa": kappa,
            "use_judge_for_headline": kappa is not None and kappa >= jcfg["kappa_min"],
            "confusion": Counter(f"{x}|{y}" for x, y in zip(m, j)),
        }
    return out


def distributions(judged: list[dict[str, Any]], labels: dict[str, list[str]]) -> dict[str, dict[str, int]]:
    valid = [j["labels"] for j in judged if j["labels"]]
    return {c: {v: sum(x[c] == v for x in valid) for v in labels[c]} for c in CRITERIA}


def judge_report(s: Setup) -> dict[str, Any]:
    jcfg = s.ecfg["judge"]
    data = load_contexts(s)
    gold = {q["id"]: set(q["gold_chunk_ids"]) for q in s.questions}
    res: dict[str, Any] = {"meta": {"generated": datetime.now().isoformat(timespec="minutes"),
                                    "rubric_version": rubric_version(jcfg)}, "answerers": {}}
    for judge, answerer in jcfg["pairs"].items():
        answers = judged_answers(s, answerer, data)
        js = current_judgements(s, judge, data)
        judged = [js[a["key"]] for a in answers if a["key"] in js]
        hit = {a["key"]: bool(gold[a["id"]] & set(data["questions"][a["id"]]["context_ids"])) for a in answers}
        by_slice = {}
        for sl in s.ecfg["slices"]:
            if sl == OUT_OF_SCOPE:
                continue
            by_slice[sl] = distributions([j for j in judged if j["slice"] == sl], jcfg["labels"])["correctness"]
        gens = [g for j in judged for g in j["generations"] if g]
        res["answerers"][answerer] = {
            "judge": judge, "judge_model": s.acfg["llm"]["profiles"][judge]["model"],
            "answers": len(answers), "judged": len(judged), "invalid": sum(j["labels"] is None for j in judged),
            "not_judged": [a["id"] for a in answers if a["key"] not in js],
            "distributions": distributions(judged, jcfg["labels"]),
            "correctness_by_slice": by_slice,
            "correctness_by_context_hit": {
                str(h): distributions([j for j in judged if hit[j["answer_key"]] == h], jcfg["labels"])["correctness"]
                for h in (True, False)},
            "judge_tokens": [sum(g["prompt_tokens"] or 0 for g in gens), sum(g["output_tokens"] or 0 for g in gens)],
            "judge_cost_usd": sum(g.get("cost_usd") or 0.0 for g in gens),
        }
    m = jcfg["manual"]
    xlsx, mapping_path = resolve(m["xlsx"]), resolve(m["sample"])
    if xlsx.exists() and mapping_path.exists():
        manual, errors = read_manual(xlsx, jcfg["labels"])
        mapping = json.loads(mapping_path.read_text(encoding="utf-8"))
        if errors:
            res["agreement"] = {"status": "incomplete", "errors": errors}
        else:
            ag = agreement(s, data, manual, mapping)
            ag["status"] = "complete"
            ag["manual_distributions"] = {
                llm: {c: dict(Counter(manual[x["sample_id"]][c] for x in mapping if x["llm"] == llm)) for c in CRITERIA}
                for llm in s.ecfg["answers"]["llms"]}
            res["agreement"] = ag
    else:
        res["agreement"] = {"status": "not exported yet (maritime-eval-judge export-manual)"}
    return res


def _dist(d: dict[str, int]) -> str:
    n = sum(d.values())
    return ", ".join(f"{k} {v}" + (f" ({v / n:.0%})" if n else "") for k, v in d.items())


def render_judge_md(res: dict[str, Any]) -> str:
    out = ["# LLM-as-judge (Phase 4, Step 4)", "",
           "Cross-judging: each LLM's fixed Step 3 answers (in-scope, answered) are graded by the other LLM. The judge "
           "sees only the question, the gold text, the reference answer, the cited sources and the answer "
           f"(`docs/eval_plan.md` section 4, post-hoc change 2). Rubric version `{res['meta']['rubric_version']}`.", ""]
    ag = res.get("agreement", {})
    usable = {c: v["use_judge_for_headline"] for c, v in ag.get("criteria", {}).items()} if ag.get("status") == "complete" else {}
    for answerer, v in res["answerers"].items():
        out += [f"## Answers of {answerer} (judged by {v['judge']}, `{v['judge_model']}`)", "",
                f"{v['judged']}/{v['answers']} judged, {v['invalid']} invalid judge outputs"
                + (f"; not judged yet: {', '.join(v['not_judged'])}" if v["not_judged"] else "")
                + f". Judge tokens {v['judge_tokens'][0]} in / {v['judge_tokens'][1]} out, est. ${v['judge_cost_usd']:.4f}.", ""]
        for c in CRITERIA:
            note = "" if c not in usable or usable[c] else " (kappa below the rule: not used for the headline)"
            out.append(f"- {c}: {_dist(v['distributions'][c])}{note}")
        out += ["", "Correctness by slice: " + "; ".join(f"{sl}: {_dist(d)}" for sl, d in v["correctness_by_slice"].items()),
                "", "Correctness by context hit: " + "; ".join(
                    f"{'gold in context' if h == 'True' else 'gold NOT in context'}: {_dist(d)}"
                    for h, d in v["correctness_by_context_hit"].items()), ""]
    out += ["## Agreement with the manual labels", ""]
    if ag.get("status") != "complete":
        out.append(f"Status: {ag.get('status')}." + (" Problems: " + "; ".join(ag["errors"]) if ag.get("errors") else ""))
    else:
        out += [f"n = {ag['n']} (kappa is unstable at this n)." + (
            f" Missing judgements: {', '.join(ag['missing_judgements'])}." if ag["missing_judgements"] else ""), "",
            "| criterion | agreement | Cohen's kappa | judge used for headline (kappa >= rule) |", "|---|---|---|---|"]
        for c, x in ag["criteria"].items():
            k = "n/a" if x["kappa"] is None else f"{x['kappa']:.2f}"
            out.append(f"| {c} | {x['agreement']:.0%} | {k} | {'yes' if x['use_judge_for_headline'] else 'NO: manual labels (n = 10 per LLM) reported instead'} |")
        out += ["", "Manual labels on the sample: " + "; ".join(
            f"{llm}: " + ", ".join(f"{c} {_dist(d)}" for c, d in dd.items()) for llm, dd in ag["manual_distributions"].items())]
    return "\n".join(out) + "\n"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("command", choices=["run", "export-manual", "report"])
    parser.add_argument("--config", default="configs/eval.yaml")
    parser.add_argument("--judge", help="run: judge profile (groq | groq-qwen); it judges the other LLM's answers")
    args = parser.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    logging.basicConfig(level=logging.WARNING, format="%(levelname)s %(name)s: %(message)s")
    s = setup(args.config)
    j = s.ecfg["judge"]
    if args.command == "run":
        if args.judge not in j["pairs"]:
            sys.exit(f"--judge must be one of {list(j['pairs'])}")
        logging.getLogger("maritime_rag.answer.providers").setLevel(logging.INFO)
        try:
            run_judge(s, args.judge)
        except ProviderError as e:
            sys.exit(f"error: {e}")
    elif args.command == "export-manual":
        xlsx, mapping = resolve(j["manual"]["xlsx"]), resolve(j["manual"]["sample"])
        if xlsx.exists():
            sys.exit(f"{xlsx} exists and may hold your labels; move it away first")
        sample = export_manual(s, load_contexts(s), xlsx, mapping)
        per = Counter((x["llm"], x["slice"]) for x in sample)
        print(f"wrote {xlsx} ({len(sample)} answers, blind) and {mapping}: " + ", ".join(f"{k[0]}/{k[1]} {v}" for k, v in sorted(per.items())))
    else:
        res = judge_report(s)
        out_json, out_md = resolve(j["output_json"]), resolve(j["output_md"])
        out_json.write_text(json.dumps(res, indent=2, ensure_ascii=False, default=dict), encoding="utf-8")
        out_md.write_text(render_judge_md(res), encoding="utf-8")
        print(f"wrote {out_json} and {out_md}; agreement: {res['agreement'].get('status')}")


if __name__ == "__main__":
    main()
