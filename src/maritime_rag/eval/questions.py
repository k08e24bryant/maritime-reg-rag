"""Eval question set: export the draft for manual review, validate the reviewed file -> gold CSV.

Usage:
  uv run maritime-eval-questions export      # data/eval/questions_draft.yaml -> questions_draft.xlsx
  uv run maritime-eval-questions validate    # reviewed questions_draft.xlsx -> questions_gold.csv

The review workbook has one row per question with a dropdown decision (keep / edit / delete),
`corrected_question` / `corrected_gold` for edits, and an `auto_check` column showing the
longest run of words the question shares with its gold chunk text (copied wording inflates BM25).
`validate` writes nothing unless every row is decided and every gold chunk id exists.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import yaml

from maritime_rag.answer.versions import DocNames, VersionIndex, page_label, status_of
from maritime_rag.config import load_config, resolve
from maritime_rag.retrieval.cli import locator
from maritime_rag.retrieval.corpus import Chunk, Corpus, load_corpus

COLUMNS = [
    "id", "slice", "question", "gold_chunk_ids", "gold chunk text", "reference_answer",
    "my_decision", "corrected_question", "corrected_gold", "notes", "auto_check",
]
GOLD_CSV_COLUMNS = ["id", "slice", "lang", "question", "gold_chunk_ids", "reference_answer", "review", "notes"]
OUT_OF_SCOPE = "out_of_scope"
EXCEL_CELL_MAX = 32_000


@dataclass
class Question:
    id: str
    slice: str
    lang: str
    question: str
    gold_chunk_ids: list[str]
    reference_answer: str


def load_draft(path: Path) -> list[Question]:
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    return [Question(q["id"], q["slice"], q["lang"], q["question"].strip(), list(q.get("gold_chunk_ids") or []),
                     q["reference_answer"].strip()) for q in data["questions"]]


def _words(text: str) -> list[str]:
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode().lower()
    return re.findall(r"[a-z0-9]+", text)


def longest_shared_run(question: str, text: str) -> list[str]:
    """Longest run of consecutive words that appears in both (case- and accent-insensitive)."""
    a, b = _words(question), _words(text)
    best_len, best_end = 0, 0
    prev = [0] * (len(b) + 1)
    for i in range(1, len(a) + 1):
        cur = [0] * (len(b) + 1)
        for j in range(1, len(b) + 1):
            if a[i - 1] == b[j - 1]:
                cur[j] = prev[j - 1] + 1
                if cur[j] > best_len:
                    best_len, best_end = cur[j], i
        prev = cur
    return a[best_end - best_len:best_end]


def parse_ids(cell: Any) -> list[str]:
    """'a, b; c' -> [a, b, c]; 'none' or '-' -> []."""
    text = str(cell or "").strip()
    if text.lower() in {"", "none", "-"}:
        return []
    return [t for t in re.split(r"[,;\s]+", text) if t]


class GoldText:
    """Renders gold chunks for review: '[id | PM 74/2016 Pasal 71 ayat (1)-(6) | p. 11-12 | CURRENT]' + text."""

    def __init__(self, corpus: Corpus) -> None:
        self.by_id = corpus.by_id()
        self.index = VersionIndex(corpus)
        self.names = DocNames(corpus)

    def status(self, chunk: Chunk) -> str:
        status, by = status_of(chunk, self.index)
        return status if by is None else f"{status} by {self.names(by)}"

    def render(self, ids: list[str]) -> str:
        parts = []
        for cid in ids:
            c = self.by_id.get(cid)
            if c is None:
                parts.append(f"[{cid} | UNKNOWN CHUNK ID]")
                continue
            head = f"[{cid} | {locator(c)} | {page_label(c)} | {self.status(c)}]"
            parts.append(f"{head}\n{' '.join(c['text'].split())}")
        return "\n\n".join(parts)[:EXCEL_CELL_MAX]


def auto_check(q: Question, gold: GoldText, warn_at: int) -> str:
    unknown = [cid for cid in q.gold_chunk_ids if cid not in gold.by_id]
    if unknown:
        return f"UNKNOWN gold ids: {', '.join(unknown)}"
    if not q.gold_chunk_ids:
        return "no gold (out of scope)" if q.slice == OUT_OF_SCOPE else "MISSING gold"
    run = max((longest_shared_run(q.question, gold.by_id[c]["text"]) for c in q.gold_chunk_ids), key=len)
    flag = "WARN copied wording? " if len(run) >= warn_at else ""
    return f"{flag}longest shared run: {len(run)} words ('{' '.join(run)}')"


def export_xlsx(questions: list[Question], corpus: Corpus, out: Path, warn_at: int) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    gold = GoldText(corpus)
    wb = Workbook()
    ws = wb.active
    ws.title = "questions"
    ws.append(COLUMNS)
    for q in questions:
        ws.append([q.id, q.slice, q.question, ", ".join(q.gold_chunk_ids), gold.render(q.gold_chunk_ids),
                   q.reference_answer, "", "", "", "", auto_check(q, gold, warn_at)])
    widths = {"id": 6, "slice": 13, "question": 50, "gold_chunk_ids": 24, "gold chunk text": 80,
              "reference_answer": 50, "my_decision": 13, "corrected_question": 40, "corrected_gold": 24,
              "notes": 30, "auto_check": 30}
    for i, name in enumerate(COLUMNS, start=1):
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = widths[name]
        ws.cell(row=1, column=i).font = Font(bold=True)
        ws.cell(row=1, column=i).fill = PatternFill("solid", fgColor="DDEBF7")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    decision_col = ws.cell(row=1, column=COLUMNS.index("my_decision") + 1).column_letter
    dv = DataValidation(type="list", formula1='"keep,edit,delete"', allow_blank=True,
                        error="Choose keep, edit or delete", errorTitle="Decision")
    ws.add_data_validation(dv)
    dv.add(f"{decision_col}2:{decision_col}{len(questions) + 1}")
    for i in range(2, len(questions) + 2):
        ws.cell(row=i, column=COLUMNS.index("my_decision") + 1).fill = PatternFill("solid", fgColor="FFF2CC")
    ws.freeze_panes = "D2"

    readme = wb.create_sheet("README")
    for line in [
        "How to review (rules: docs/eval_plan.md, section 1)",
        "",
        "my_decision (dropdown): keep | edit | delete. Every row needs one.",
        "edit: fill corrected_question and/or corrected_gold; an empty corrected_* cell keeps the original.",
        "corrected_gold: chunk ids separated by commas (e.g. pm_74_2016-0010, pm_93_2013-0084); 'none' = no gold.",
        "reference_answer can be edited in place; validate takes the cell as it is.",
        "Gold rules: every chunk that contains (part of) the answer; amended slice = current version + "
        "superseded base; out_of_scope = no gold.",
        "auto_check: longest run of consecutive words shared by the question and its gold text. WARN at >= "
        f"{warn_at} words: rephrase if it is copied wording rather than a legal term.",
        "gold chunk text shows each chunk as [id | location | pages | version status].",
        "When done: save this file and run `uv run maritime-eval-questions validate`.",
    ]:
        readme.append([line])
    readme.column_dimensions["A"].width = 120
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


@dataclass
class ValidationResult:
    rows: list[dict[str, str]]
    errors: list[str]
    warnings: list[str]


def validate_reviewed(xlsx: Path, draft: list[Question], corpus: Corpus, cfg: dict[str, Any],
                      detect: Any = None) -> ValidationResult:
    """Apply the review decisions. detect(text) -> 'id' | 'en' re-detects edited questions' language."""
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb["questions"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    errors, warnings, out = [], [], []
    missing_cols = [c for c in COLUMNS[:10] if c not in header]
    if missing_cols:
        return ValidationResult([], [f"missing columns: {', '.join(missing_cols)}"], [])
    col = {name: header.index(name) for name in header if name}
    by_id = {q.id: q for q in draft}
    known = corpus.by_id()
    decisions = set(cfg["review"]["decisions"])
    seen: set[str] = set()

    def cell(r: tuple, name: str) -> str:
        v = r[col[name]] if col[name] < len(r) else None
        return str(v).strip() if v is not None else ""

    for n, r in enumerate(rows[1:], start=2):
        qid = cell(r, "id")
        if not qid:
            continue
        where = f"row {n} ({qid})"
        if qid not in by_id:
            errors.append(f"{where}: id not in the draft")
            continue
        if qid in seen:
            errors.append(f"{where}: duplicate id")
            continue
        seen.add(qid)
        decision = cell(r, "my_decision").lower()
        if decision not in decisions:
            errors.append(f"{where}: my_decision must be one of {sorted(decisions)} (got {decision or 'empty'!r})")
            continue
        if decision == "delete":
            continue
        q = by_id[qid]
        slice_ = cell(r, "slice") or q.slice
        new_q, new_gold = cell(r, "corrected_question"), cell(r, "corrected_gold")
        if decision == "edit" and not (new_q or new_gold):
            errors.append(f"{where}: decision is edit but corrected_question and corrected_gold are empty")
            continue
        question = new_q or cell(r, "question") or q.question
        gold = parse_ids(new_gold) if new_gold else parse_ids(cell(r, "gold_chunk_ids"))
        unknown = [g for g in gold if g not in known]
        if unknown:
            errors.append(f"{where}: unknown gold chunk ids {unknown}")
        if slice_ == OUT_OF_SCOPE and gold:
            errors.append(f"{where}: out_of_scope must have no gold chunk ids")
        if slice_ != OUT_OF_SCOPE and not gold:
            errors.append(f"{where}: in-scope question needs at least one gold chunk id")
        if slice_ not in cfg["slices"]:
            errors.append(f"{where}: unknown slice {slice_!r}")
        reference = cell(r, "reference_answer")
        if slice_ != OUT_OF_SCOPE and not reference:
            warnings.append(f"{where}: empty reference_answer")
        lang = detect(question) if (new_q and detect) else q.lang
        out.append({"id": qid, "slice": slice_, "lang": lang, "question": question,
                    "gold_chunk_ids": ";".join(gold), "reference_answer": reference, "review": decision,
                    "notes": cell(r, "notes")})

    not_reviewed = [q.id for q in draft if q.id not in seen]
    if not_reviewed:
        errors.append(f"rows missing from the workbook: {', '.join(not_reviewed)}")
    texts = [o["question"].lower() for o in out]
    dupes = sorted({t for t in texts if texts.count(t) > 1})
    if dupes:
        errors.append(f"duplicate question text: {dupes}")
    for slice_, target in cfg["slices"].items():
        kept = sum(1 for o in out if o["slice"] == slice_)
        if kept < 0.8 * target:
            warnings.append(f"slice {slice_}: {kept} questions kept, below 80 % of the planned {target}")
    return ValidationResult(out, errors, warnings)


def write_gold_csv(rows: list[dict[str, str]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=GOLD_CSV_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)


def load_gold_csv(path: Path) -> list[dict[str, Any]]:
    with path.open(encoding="utf-8", newline="") as f:
        rows = list(csv.DictReader(f))
    for r in rows:
        r["gold_chunk_ids"] = [g for g in r["gold_chunk_ids"].split(";") if g]
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("command", choices=["export", "validate"])
    parser.add_argument("--config", default="configs/eval.yaml")
    parser.add_argument("--xlsx", help="review workbook (default: paths.draft_xlsx)")
    args = parser.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")

    cfg = load_config(args.config)
    p = cfg["paths"]
    corpus = load_corpus(resolve(p["chunks_file"]))
    draft = load_draft(resolve(p["draft"]))
    xlsx = resolve(args.xlsx or p["draft_xlsx"])

    if args.command == "export":
        if xlsx.exists():
            sys.exit(f"{xlsx} exists; it may hold your review. Move it away first, or pass --xlsx NEW_PATH.")
        export_xlsx(draft, corpus, xlsx, cfg["review"]["shared_words_warn"])
        counts = {s: sum(1 for q in draft if q.slice == s) for s in cfg["slices"]}
        print(f"wrote {xlsx}: {len(draft)} questions {counts}")
        return

    from maritime_rag.answer.prompt import LanguageDetector

    rcfg = load_config("configs/retrieval.yaml")
    detect = LanguageDetector(resolve(rcfg["paths"]["stopwords"]))
    result = validate_reviewed(xlsx, draft, corpus, cfg, detect)
    for w in result.warnings:
        print(f"warning: {w}")
    if result.errors:
        for e in result.errors:
            print(f"error: {e}", file=sys.stderr)
        sys.exit(f"{len(result.errors)} error(s); nothing written")
    out = resolve(p["gold_csv"])
    write_gold_csv(result.rows, out)
    counts = {s: sum(1 for r in result.rows if r["slice"] == s) for s in cfg["slices"]}
    edits = sum(1 for r in result.rows if r["review"] == "edit")
    print(f"wrote {out}: {len(result.rows)} questions ({edits} edited, {len(draft) - len(result.rows)} deleted) {counts}")


if __name__ == "__main__":
    main()
