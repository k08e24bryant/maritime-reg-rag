"""Phase 4 question set: overlap check, review workbook export, validation of the reviewed file."""

from __future__ import annotations

import pytest

from maritime_rag.config import load_config, resolve
from maritime_rag.eval.questions import (
    COLUMNS,
    GoldText,
    Question,
    auto_check,
    export_xlsx,
    load_draft,
    load_gold_csv,
    longest_shared_run,
    parse_ids,
    validate_reviewed,
    write_gold_csv,
)
from maritime_rag.retrieval.corpus import Corpus, load_corpus

CFG = load_config("configs/eval.yaml")


def _chunk(cid: str, pasal: str, text: str, **kw) -> dict:
    doc = cid.rsplit("-", 1)[0]
    return {"chunk_id": cid, "doc_id": doc, "doc_type": "PM", "number": doc.split("_")[1], "year": 2013,
            "title": "T", "part": "body", "pasal": pasal, "ayat": None, "heading": f"{doc} > Pasal {pasal}",
            "source_file": f"{doc}.pdf", "page_start": 4, "page_end": 4, "text": text, "section": None,
            "subsection": None, "amends_doc": None, "amends_pasal": None, "amendment_kind": None,
            "superseded_by": [], "partially_amended_by": [], "may_be_amended_by_out_of_corpus": []} | kw


@pytest.fixture
def corpus() -> Corpus:
    return Corpus([
        _chunk("pm_93_2013-0001", "5", "Pasal 5 (1) Izin usaha angkutan laut diberikan oleh Direktur Jenderal.",
               superseded_by=["pm_74_2016"]),
        _chunk("pm_74_2016-0001", "5", "Pasal 5 (1) Izin usaha angkutan laut diberikan melalui OSS.",
               amends_doc="pm_93_2013", amends_pasal="5", amendment_kind="replace", year=2016),
        _chunk("pm_93_2013-0002", "9", "Pasal 9 Kapal wajib membawa surat ukur."),
    ], "sha")


DRAFT = [
    Question("A01", "amended", "id", "Siapa yang menerbitkan izin pelayaran?", ["pm_74_2016-0001", "pm_93_2013-0001"],
             "Melalui OSS (PM 74/2016)."),
    Question("F01", "factual_id", "id", "Dokumen apa yang harus dibawa kapal?", ["pm_93_2013-0002"], "Surat ukur."),
    Question("O01", "out_of_scope", "id", "Berapa harga tiket pesawat?", [], "Tidak ada."),
]


def test_longest_shared_run_and_parse_ids() -> None:
    assert longest_shared_run("Kapan izin usaha angkutan laut diberikan?",
                              "Izin usaha angkutan laut diberikan oleh Menteri") == ["izin", "usaha", "angkutan", "laut", "diberikan"]
    assert longest_shared_run("Café?", "cafe") == ["cafe"]
    assert parse_ids("a-1, b-2; c-3  d-4") == ["a-1", "b-2", "c-3", "d-4"]
    assert parse_ids("none") == [] and parse_ids("-") == [] and parse_ids(None) == []


def test_auto_check_and_gold_text(corpus: Corpus) -> None:
    gold = GoldText(corpus)
    assert auto_check(DRAFT[0], gold, 5).startswith("longest shared run")
    copied = Question("X", "factual_id", "id", "Izin usaha angkutan laut diberikan oleh siapa?", ["pm_93_2013-0001"], "-")
    assert auto_check(copied, gold, 5).startswith("WARN copied wording? longest shared run: 6 words")
    assert auto_check(Question("Y", "factual_id", "id", "q", ["nope-1"], "-"), gold, 5) == "UNKNOWN gold ids: nope-1"
    text = gold.render(["pm_93_2013-0001"])
    assert text.startswith("[pm_93_2013-0001 | PM 93/2013 Pasal 5 | p. 4 | SUPERSEDED by PM 74/2016]")


def _review(path, decisions: dict[str, dict]) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb["questions"]
    col = {name: i + 1 for i, name in enumerate(COLUMNS)}
    for row in range(2, ws.max_row + 1):
        qid = ws.cell(row=row, column=col["id"]).value
        for name, value in decisions.get(qid, {}).items():
            ws.cell(row=row, column=col[name]).value = value
    wb.save(path)


def test_export_has_dropdown_and_readme(corpus: Corpus, tmp_path) -> None:
    from openpyxl import load_workbook

    out = tmp_path / "draft.xlsx"
    export_xlsx(DRAFT, corpus, out, 5)
    wb = load_workbook(out)
    ws = wb["questions"]
    assert [c.value for c in ws[1]] == COLUMNS
    assert ws.max_row == 4 and "README" in wb.sheetnames
    dvs = ws.data_validations.dataValidation
    assert len(dvs) == 1 and dvs[0].formula1 == '"keep,edit,delete"' and "G2:G4" in str(dvs[0].sqref)


def test_validate_applies_decisions(corpus: Corpus, tmp_path) -> None:
    out = tmp_path / "draft.xlsx"
    export_xlsx(DRAFT, corpus, out, 5)
    _review(out, {
        "A01": {"my_decision": "keep"},
        "F01": {"my_decision": "edit", "corrected_question": "Surat apa yang wajib ada di kapal?",
                "notes": "rephrased"},
        "O01": {"my_decision": "delete"},
    })
    result = validate_reviewed(out, DRAFT, corpus, CFG, detect=lambda t: "id")
    assert result.errors == []
    assert [(r["id"], r["review"]) for r in result.rows] == [("A01", "keep"), ("F01", "edit")]
    assert result.rows[0]["gold_chunk_ids"] == "pm_74_2016-0001;pm_93_2013-0001"
    assert result.rows[1]["question"] == "Surat apa yang wajib ada di kapal?" and result.rows[1]["notes"] == "rephrased"
    csv_path = tmp_path / "gold.csv"
    write_gold_csv(result.rows, csv_path)
    rows = load_gold_csv(csv_path)
    assert rows[0]["gold_chunk_ids"] == ["pm_74_2016-0001", "pm_93_2013-0001"]


def test_validate_reports_errors_and_writes_nothing(corpus: Corpus, tmp_path) -> None:
    out = tmp_path / "draft.xlsx"
    export_xlsx(DRAFT, corpus, out, 5)
    _review(out, {
        "A01": {"my_decision": "edit"},                                      # edit without corrections
        "F01": {"my_decision": "keep", "gold_chunk_ids": "pm_93_2013-9999"},  # unknown gold
        "O01": {"my_decision": "edit", "corrected_gold": "pm_93_2013-0002"},  # out of scope with gold
    })
    errors = validate_reviewed(out, DRAFT, corpus, CFG).errors
    assert any("A01" in e and "edit but" in e for e in errors)
    assert any("F01" in e and "unknown gold" in e for e in errors)
    assert any("O01" in e and "out_of_scope must have no gold" in e for e in errors)
    _review(out, {"A01": {"my_decision": None}})
    assert any("A01" in e and "empty" in e for e in validate_reviewed(out, DRAFT, corpus, CFG).errors)


# --- the real draft (needs chunks.jsonl) ----------------------------------------------------------

CHUNKS = resolve(CFG["paths"]["chunks_file"])


@pytest.mark.skipif(not CHUNKS.exists(), reason="needs chunks.jsonl (maritime-ingest)")
def test_real_draft_is_consistent() -> None:
    corpus = load_corpus(CHUNKS)
    known = corpus.by_id()
    draft = load_draft(resolve(CFG["paths"]["draft"]))
    assert len({q.id for q in draft}) == len(draft)
    assert {s: sum(q.slice == s for q in draft) for s in CFG["slices"]} == CFG["slices"]
    for q in draft:
        assert all(g in known for g in q.gold_chunk_ids), q.id
        assert bool(q.gold_chunk_ids) == (q.slice != "out_of_scope"), q.id
        assert q.lang == ("en" if q.slice == "english" or q.id in {"O09", "O10"} else "id"), q.id
    gold = GoldText(corpus)
    assert not [q.id for q in draft if auto_check(q, gold, CFG["review"]["shared_words_warn"]).startswith("WARN")]
    # amended slice: every question has a current amending chunk and a superseded base chunk as gold
    for q in (q for q in draft if q.slice == "amended"):
        chunks = [known[g] for g in q.gold_chunk_ids]
        assert any(c.get("amends_doc") and not c.get("superseded_by") for c in chunks), q.id
        assert any(c.get("superseded_by") for c in chunks), q.id
