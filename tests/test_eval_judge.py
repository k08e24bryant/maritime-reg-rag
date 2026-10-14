"""Phase 4 Step 4 with fake LLMs: judge input, parsing, caching/resume, blind sample, kappa, agreement."""

from __future__ import annotations

import io
import json

import pytest
from test_eval_answers import FakeLLM, s  # noqa: F401  (fixture: stored contexts for A01, F01, O01)

from maritime_rag.answer.providers import Generation, ProviderError
from maritime_rag.eval.answers import Setup, load_contexts, run_llm
from maritime_rag.eval.judge import (
    MANUAL_COLUMNS,
    cohen_kappa,
    current_judgements,
    export_manual,
    judge_messages,
    judge_report,
    judged_answers,
    largest_remainder,
    parse_judgement,
    read_manual,
    render_judge_md,
    run_judge,
    stratified_sample,
)
from maritime_rag.eval.questions import GoldText

GOOD = '{"correctness": "correct", "faithfulness": "supported", "superseded_mentioned": "n.a.", "rationale": "ok"}'


class FakeJudge:
    name, model = "judge", "fake-judge"

    def __init__(self, replies: list[str] | None = None, fail_at: int | None = None) -> None:
        self.replies, self.fail_at, self.n, self.seen = list(replies or []), fail_at, 0, []
        self.calls = self.prompt_tokens = self.output_tokens = 0
        self.cost_usd = 0.0

    def check(self) -> None:
        pass

    def generate(self, messages):
        self.n += 1
        if self.fail_at == self.n:
            raise ProviderError("Groq returned HTTP 429 after 5 retries: rate limit")
        self.seen.append(messages)
        self.calls += 1
        text = self.replies.pop(0) if self.replies else GOOD
        return Generation(text, "judge", self.model, 800, 60, 0.3, finish_reason="stop", cost_usd=0.0002)


def _answers(s: Setup) -> None:  # noqa: F811
    run_llm(s, "groq", FakeLLM(), io.StringIO())
    run_llm(s, "groq-qwen", FakeLLM(), io.StringIO())


def test_parse_judgement_variants() -> None:
    labels = {"correctness": ["correct", "partial", "wrong"], "faithfulness": ["supported", "partly", "unsupported"],
              "superseded_mentioned": ["yes", "no", "n.a."]}
    got, err = parse_judgement("```json\n" + GOOD.replace('"n.a."', '"N/A"') + "\n```", labels)
    assert err is None and got["superseded_mentioned"] == "n.a." and got["rationale"] == "ok"
    assert parse_judgement("no json here", labels) == (None, "no JSON object in the judge output")
    got, err = parse_judgement(GOOD.replace('"correct"', '"mostly right"'), labels)
    assert got is None and "correctness" in err


def test_judge_sees_cited_sources_only(s: Setup) -> None:  # noqa: F811
    _answers(s)
    data = load_contexts(s)
    a01 = next(a for a in judged_answers(s, "groq-qwen", data) if a["id"] == "A01")
    user = judge_messages(s, GoldText(s.corpus), a01)[1]["content"]
    cited = user.split("CITED SOURCES")[1].split("ANSWER:")[0]
    assert "pm_74_2016-0001" in cited and "pm_93_2013-0001" not in cited  # in the context, not cited
    assert "SUPERSEDED by PM 74/2016" in user.split("CITED SOURCES")[0]  # gold carries its status
    assert user.rstrip().endswith(a01["text"].rstrip())


def test_run_judge_caches_retries_and_resumes(s: Setup) -> None:  # noqa: F811
    _answers(s)
    out = io.StringIO()
    judge = FakeJudge(replies=["not json", GOOD], fail_at=3)  # A01: invalid then valid; F01: quota stop
    assert run_judge(s, "groq", judge, out) == 1
    assert "stopped (quota / spend limit) after 1 new judgements" in out.getvalue()
    data = load_contexts(s)
    js = current_judgements(s, "groq", data)
    assert len(js) == 1 and next(iter(js.values()))["labels"]["correctness"] == "correct"
    assert len(next(iter(js.values()))["generations"]) == 2  # the retry is recorded
    judge = FakeJudge()
    assert run_judge(s, "groq", judge, io.StringIO()) == 1  # only F01 left; O01 (refused) is never judged
    assert run_judge(s, "groq", FakeJudge(), io.StringIO()) == 0


def test_largest_remainder_and_kappa() -> None:
    assert largest_remainder({"factual_id": 24, "which_pasal": 7, "amended": 8, "english": 5}, 10) == {
        "factual_id": 5, "which_pasal": 2, "amended": 2, "english": 1}
    assert sum(largest_remainder({"a": 3, "b": 3, "c": 3}, 10).values()) == 9  # capped by availability
    assert cohen_kappa(["x", "x", "y", "y"], ["x", "y", "y", "y"]) == pytest.approx(0.5)
    assert cohen_kappa(["x", "y"], ["x", "y"]) == 1.0 and cohen_kappa([], []) is None
    assert cohen_kappa(["x", "x"], ["x", "x"]) == 1.0  # one category only: agreement 1, kappa defined as 1


def test_blind_export_and_agreement(s: Setup, tmp_path) -> None:  # noqa: F811
    from openpyxl import load_workbook

    _answers(s)
    run_judge(s, "groq", FakeJudge(), io.StringIO())
    run_judge(s, "groq-qwen", FakeJudge(), io.StringIO())
    data = load_contexts(s)
    sample = stratified_sample(s, data)
    assert [x["sample_id"] for x in sample] == ["M01", "M02", "M03", "M04"]  # 2 judged answers per LLM here
    assert sample == stratified_sample(s, data)  # deterministic
    xlsx, mapping = tmp_path / "manual.xlsx", tmp_path / "sample.json"
    export_manual(s, data, xlsx, mapping)
    wb = load_workbook(xlsx)
    ws = wb["labels"]
    header = [c.value for c in ws[1]]
    assert header == MANUAL_COLUMNS and "llm" not in header and "groq" not in str([c.value for c in ws[2]])
    assert len(ws.data_validations.dataValidation) == 3
    for r in range(2, ws.max_row + 1):  # label like the fake judge, except one disagreement
        ws.cell(r, header.index("correctness") + 1).value = "partial" if r == 2 else "correct"
        ws.cell(r, header.index("faithfulness") + 1).value = "supported"
        ws.cell(r, header.index("superseded_mentioned") + 1).value = "n.a."
    wb.save(xlsx)
    manual, errors = read_manual(xlsx, s.ecfg["judge"]["labels"])
    assert errors == [] and len(manual) == 4
    s.ecfg["judge"]["manual"] |= {"xlsx": str(xlsx), "sample": str(mapping)}
    res = judge_report(s)
    ag = res["agreement"]
    assert ag["status"] == "complete" and ag["n"] == 4
    assert ag["criteria"]["correctness"]["agreement"] == 0.75 and ag["criteria"]["faithfulness"]["kappa"] == 1.0
    assert not ag["criteria"]["correctness"]["use_judge_for_headline"]  # kappa 0 < 0.40
    md = render_judge_md(res)
    assert "| correctness | 75% | 0.00 | NO: manual labels" in md and "kappa below the rule" in md
    assert res["answerers"]["groq-qwen"]["judge"] == "groq" and res["answerers"]["groq"]["judge"] == "groq-qwen"
